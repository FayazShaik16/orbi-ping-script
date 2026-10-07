#!/usr/bin/env python3
"""
Rebuild Master Excel Workbook from CSV Journals
------------------------------------------------
Reads all 12 hours of ping journals from ping_logs/ and constructs
the complete 6-sheet Master Excel Workbook with:
- Executive 2-row summary header
- Hourly divider banners
- All data rows
- Hourly breakdown tables
- Clean Bar Chart at H1 (top of sheet beside stats, no Y-axis labels, no gridlines)
"""

import os
import sys
import csv
import time
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

SHEET_CONFIGS = [
    {"name": "Base_Internet", "node_type": "Base", "title": "Base Node Client -> Internet (8.8.8.8)", "file": "Base_Internet.csv"},
    {"name": "Base_Self", "node_type": "Base", "title": "Base Node Client -> Base Node IP (10.168.168.1)", "file": "Base_Self.csv"},
    {"name": "Base_Satellite1", "node_type": "Base", "title": "Base Node Client -> Satellite-1 IP", "file": "Base_Satellite1.csv"},
    {"name": "Base_Satellite2", "node_type": "Base", "title": "Base Node Client -> Satellite-2 IP", "file": "Base_Satellite2.csv"},
    {"name": "Satellite1_Internet", "node_type": "Satellite1", "title": "Satellite-1 (SSH) -> Internet (8.8.8.8)", "file": "Satellite1_Internet.csv"},
    {"name": "Satellite2_Internet", "node_type": "Satellite2", "title": "Satellite-2 (SSH) -> Internet (8.8.8.8)", "file": "Satellite2_Internet.csv"},
]

HEADER_COLUMNS = ["Time stamp", "Hour", "Mesh Node Type", "PacketLoss", "TTL", "String"]


def rebuild_workbook(log_dir, output_paths):
    print("=" * 70)
    print(" REBUILDING 12-HOUR MASTER EXCEL WORKBOOK FROM CSV JOURNALS")
    print("=" * 70)
    start_time = time.time()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    label_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    val_font = Font(name="Calibri", size=11, bold=True)
    green_font = Font(name="Calibri", size=11, bold=True, color="27AE60")
    red_font = Font(name="Calibri", size=11, bold=True, color="C0392B")
    loss_row_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    loss_row_font = Font(name="Calibri", size=10, bold=True, color="9C0006")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    sep_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sep_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    sep_border = Border(
        left=Side(style="thin", color="1B365D"),
        right=Side(style="thin", color="1B365D"),
        top=Side(style="medium", color="0D1E3A"),
        bottom=Side(style="medium", color="0D1E3A"),
    )

    align_center = Alignment(horizontal="center", vertical="center")

    for cfg in SHEET_CONFIGS:
        s_name = cfg["name"]
        csv_file = os.path.join(log_dir, cfg["file"])
        if not os.path.exists(csv_file):
            print(f"[!] Warning: {csv_file} not found, skipping...")
            continue

        sheet_start = time.time()
        print(f"\n[*] Processing sheet '{s_name}' from {cfg['file']}...")

        ws = wb.create_sheet(title=s_name)

        # Row 1: SheetName (TestCase)
        ws["A1"] = "SheetName (TestCase):"
        ws["B1"] = f"{s_name} ({cfg['title']})"
        ws["A1"].font = label_font
        ws["B1"].font = val_font

        # Row 2: Packet Loss Stats placeholders
        ws["A2"] = "Packet Loss :"
        ws["C2"] = "Packet Loss Count :"
        ws["E2"] = "Packet Loss Hour(s) :"
        ws["A2"].font = label_font
        ws["C2"].font = label_font
        ws["E2"].font = label_font

        # Row 3: Blank separator
        ws.append([])

        # Row 4: Column Headers
        ws.append(HEADER_COLUMNS)
        for col_idx in range(1, len(HEADER_COLUMNS) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [2, 3, 4, 5] else "left")

        # Column widths & freeze panes
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 22
        ws.column_dimensions["F"].width = 65
        ws.column_dimensions["G"].width = 3
        ws.freeze_panes = "A5"

        # Read CSV data and stream rows
        total_pings = 0
        total_losses = 0
        loss_hours = set()
        hourly_stats = {}  # hour_idx -> {"total": 0, "losses": 0, "start": dt, "end": dt}

        t0 = None
        current_hour_idx = -1

        with open(csv_file, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            header = next(reader, None)

            for row in reader:
                if not row or len(row) < 3:
                    continue

                ts_val = row[0].strip()
                node_val = row[1].strip()
                loss_val = row[2].strip()
                if len(row) >= 5:
                    ttl_val = row[3].strip()
                    str_val = row[4]
                else:
                    ttl_val = "NA"
                    str_val = row[3] if len(row) > 3 else ""

                try:
                    ttl_out = int(ttl_val)
                except (ValueError, TypeError):
                    ttl_out = ttl_val

                # Parse packet loss
                loss_clean = loss_val.upper()
                if loss_clean in ["NA", "", "0"]:
                    loss_out = 0
                else:
                    try:
                        loss_out = int(loss_clean)
                    except ValueError:
                        loss_out = loss_clean

                is_loss = (loss_out != 0 and loss_clean != "0")
                total_pings += 1
                if is_loss:
                    total_losses += 1

                # Calculate hour index
                try:
                    t = datetime.strptime(ts_val.split(".")[0], "%Y-%m-%d %H:%M:%S")
                except Exception:
                    t = None

                if t0 is None and t is not None:
                    t0 = t

                if t0 and t:
                    diff_sec = max(0.0, (t - t0).total_seconds())
                    h_idx = int(diff_sec // 3600)
                else:
                    h_idx = 0

                if is_loss:
                    loss_hours.add(h_idx)

                # Hourly bucket tracking
                if h_idx not in hourly_stats:
                    h_start = (t0 + timedelta(hours=h_idx)) if t0 else t
                    h_end = (t0 + timedelta(hours=h_idx + 1)) if t0 else t
                    hourly_stats[h_idx] = {"total": 0, "losses": 0, "start": h_start, "end": h_end}

                hourly_stats[h_idx]["total"] += 1
                if is_loss:
                    hourly_stats[h_idx]["losses"] += 1

                hour_label = f"Hour {h_idx + 1}"

                # Insert Hourly Separator Banner
                if h_idx > current_hour_idx:
                    current_hour_idx = h_idx
                    h_info = hourly_stats[h_idx]
                    s_str = h_info["start"].strftime("%H:%M:%S") if h_info["start"] else ""
                    e_str = h_info["end"].strftime("%H:%M:%S") if h_info["end"] else ""
                    sep_title = f"=== HOUR {h_idx + 1} ({s_str} - {e_str}) ==="
                    sep_row = [
                        sep_title,
                        hour_label,
                        "---",
                        "---",
                        "---",
                        f"=== HOUR {h_idx + 1} START ==="
                    ]
                    ws.append(sep_row)
                    sep_row_idx = ws.max_row
                    ws.row_dimensions[sep_row_idx].height = 22
                    for c_idx in range(1, len(HEADER_COLUMNS) + 1):
                        cell = ws.cell(row=sep_row_idx, column=c_idx)
                        cell.fill = sep_fill
                        cell.font = sep_font
                        cell.alignment = align_center
                        cell.border = sep_border

                # Append data row
                ws.append([ts_val, hour_label, node_val, loss_out, ttl_out, str_val])
                if is_loss:
                    loss_r_idx = ws.max_row
                    for c_idx in range(1, len(HEADER_COLUMNS) + 1):
                        c_cell = ws.cell(row=loss_r_idx, column=c_idx)
                        c_cell.fill = loss_row_fill
                        c_cell.font = loss_row_font

        # AutoFilter over all data rows
        data_end_row = ws.max_row
        ws.auto_filter.ref = f"A4:F{data_end_row}"

        # Populate Row 2 executive stats
        loss_pct = (total_losses / total_pings * 100.0) if total_pings > 0 else 0.0
        ws["B2"] = f"{loss_pct:.2f}%"
        ws["D2"] = total_losses
        ws["B2"].font = red_font if total_losses > 0 else green_font
        ws["D2"].font = red_font if total_losses > 0 else val_font

        if not loss_hours:
            ws["F2"] = "None"
            ws["F2"].font = green_font
        else:
            loss_hrs_sorted = sorted(loss_hours)
            desc_items = []
            for h in loss_hrs_sorted:
                if t0:
                    s_t = (t0 + timedelta(hours=h)).strftime("%H:%M")
                    e_t = (t0 + timedelta(hours=h + 1)).strftime("%H:%M")
                    desc_items.append(f"Hour {h + 1} ({s_t} - {e_t})")
                else:
                    desc_items.append(f"Hour {h + 1}")
            full_str = ", ".join(desc_items)
            ws["F2"] = full_str if len(full_str) <= 60 else ", ".join(f"Hour {h + 1}" for h in loss_hrs_sorted)
            ws["F2"].font = red_font

        # Collect only the hours with packet losses
        loss_h_indices = [h for h in sorted(hourly_stats.keys()) if hourly_stats[h]["losses"] > 0]

        # Append Hourly Packet Loss Breakdown Table (Loss Hours Only) - Feeds the Top Chart
        ws.append([])
        ws.append(["--- HOURLY PACKET LOSS BREAKDOWN (LOSS HOURS ONLY) ---", "", "", ""])
        t_row = ws.max_row
        ws.cell(row=t_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

        table_header = ["Hour Window", "Total Pings", "Loss Count", "Avg Packet Loss (%)"]
        ws.append(table_header)
        header_row = ws.max_row
        for col_idx in range(1, 5):
            c = ws.cell(row=header_row, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center if col_idx in [2, 3, 4] else Alignment(horizontal="left")

        data_start_row = header_row + 1
        if not loss_h_indices:
            # Sheet had 0 losses across all hours
            ws.append(["All Hours (0% Loss)", total_pings, 0, 0.0])
            curr_row = ws.max_row
            ws.cell(row=curr_row, column=2).alignment = align_center
            ws.cell(row=curr_row, column=3).alignment = align_center
            c_pct = ws.cell(row=curr_row, column=4)
            c_pct.alignment = align_center
            c_pct.font = green_font
        else:
            for h_idx in loss_h_indices:
                st = hourly_stats[h_idx]
                s_t = st["start"].strftime("%H:%M") if st["start"] else ""
                lbl = f"Hour {h_idx + 1} ({s_t})" if s_t else f"Hour {h_idx + 1}"
                h_tot = st["total"]
                h_loss = st["losses"]
                h_pct = (h_loss / h_tot * 100.0) if h_tot > 0 else 0.0

                ws.append([lbl, h_tot, h_loss, round(h_pct, 2)])
                curr_row = ws.max_row
                ws.cell(row=curr_row, column=2).alignment = align_center
                ws.cell(row=curr_row, column=3).alignment = align_center
                c_pct = ws.cell(row=curr_row, column=4)
                c_pct.alignment = align_center
                c_pct.font = red_font if h_loss > 0 else green_font

        summary_end_row = ws.max_row

        # Build Bar Chart anchored at H1 (top of sheet beside stats)
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        if not loss_h_indices:
            chart.title = f"Average Packet Loss per 1 Hour - {s_name} (0% Loss)"
        else:
            chart.title = f"Average Packet Loss per 1 Hour - {s_name} (Loss Hours Only)"
        chart.x_axis.title = "1-Hour Time Window"
        chart.y_axis.title = None
        chart.legend = None
        chart.width = 24
        chart.height = 12

        # Fixed 0% - 100% scale
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 100
        # Remove Y-axis labels
        chart.y_axis.delete = True

        # Hide all gridlines totally
        chart.y_axis.majorGridlines = None
        chart.y_axis.minorGridlines = None
        chart.x_axis.majorGridlines = None
        chart.x_axis.minorGridlines = None

        # Ensure all category labels and ticks are rendered without skipping
        chart.x_axis.tickLblSkip = 1
        chart.x_axis.tickMarkSkip = 1

        # Direct value labels on bars
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True

        data_ref = Reference(ws, min_col=4, min_row=header_row, max_row=summary_end_row)
        cats_ref = Reference(ws, min_col=1, min_row=data_start_row, max_row=summary_end_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        ws.add_chart(chart, "H1")

        # Full 12-Hour Benchmark Table (All Hours) appended below for complete reference
        ws.append([])
        ws.append(["--- COMPLETE 12-HOUR BENCHMARK LOG (ALL HOURS) ---", "", "", ""])
        t2_row = ws.max_row
        ws.cell(row=t2_row, column=1).font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

        table2_header = ["Hour Window", "Total Pings", "Loss Count", "Avg Packet Loss (%)"]
        ws.append(table2_header)
        header2_row = ws.max_row
        for col_idx in range(1, 5):
            c = ws.cell(row=header2_row, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = align_center if col_idx in [2, 3, 4] else Alignment(horizontal="left")

        for h_idx in sorted(hourly_stats.keys()):
            st = hourly_stats[h_idx]
            s_t = st["start"].strftime("%H:%M") if st["start"] else ""
            e_t = st["end"].strftime("%H:%M") if st["end"] else ""
            lbl = f"Hour {h_idx + 1} ({s_t} - {e_t})" if s_t else f"Hour {h_idx + 1}"
            h_tot = st["total"]
            h_loss = st["losses"]
            h_pct = (h_loss / h_tot * 100.0) if h_tot > 0 else 0.0

            ws.append([lbl, h_tot, h_loss, round(h_pct, 2)])
            curr_row = ws.max_row
            ws.cell(row=curr_row, column=2).alignment = align_center
            ws.cell(row=curr_row, column=3).alignment = align_center
            c_pct = ws.cell(row=curr_row, column=4)
            c_pct.alignment = align_center
            c_pct.font = red_font if h_loss > 0 else green_font

        sheet_elapsed = time.time() - sheet_start
        print(f"  [✓] '{s_name}': {total_pings:,} pings | {total_losses} losses ({loss_pct:.2f}%) | {len(hourly_stats)} hours processed in {sheet_elapsed:.1f}s")

    print("\n[*] Saving rebuilt Master Excel Workbook to disk...")
    save_start = time.time()
    for out_path in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        wb.save(out_path)
        print(f"  [✓] Successfully saved to: {out_path} ({os.path.getsize(out_path) / (1024*1024):.1f} MB)")

    total_elapsed = time.time() - start_time
    print(f"\n[+] Master Excel Workbook rebuilt successfully in {total_elapsed:.1f}s!")
    print("=" * 70)


if __name__ == "__main__":
    log_dir = sys.argv[1] if len(sys.argv) > 1 else "/Users/fayazshaik/Documents/Candela/Script-Ping/ping_logs"
    out_paths = [
        "/Users/fayazshaik/Downloads/orbi_mesh_ping_20261006_185424.xlsx",
        os.path.join(log_dir, "orbi_mesh_ping_20261006_185424.xlsx"),
    ]
    rebuild_workbook(log_dir, out_paths)
